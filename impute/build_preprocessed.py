# -*- coding: utf-8 -*-
"""
[전처리] 로데이터 취합본 -> 결측 보완 + 가용시점 정렬 -> 데이터취합_전처리_<YYYYMMDD>.xlsx

로데이터 취합본(merge/build_panel.py)과 같은 입력(processed/)에서 같은 패널을 만든 뒤,
검토 문서(docs/결측치_처리방안_검토요청)의 ★안대로 빈칸을 보완한다. 로데이터 취합본 파일은
건드리지 않는다.

시트
  1_개요 / 2_데이터목록 / 3_전처리방법 / 4_공표시차규칙   설명서 (데이터 시트에는 꼬리표 없음)
  17시도_1차_결측보완              1단계. 로데이터 취합본의 빈칸을 보완 (기준기간 기준) + 꼬리표
  서울25구_1차_결측보완            위와 같음 + 서울 상위지역 변수
  17시도_2차_공표시점반영(ML용)     2단계. 1차 결과를 공표 시점만큼 밀어 '그 달 말에 알 수 있던 값'만 둠
                                   (+ 분기·반기·연간을 공표월부터 월별로 채움)
  서울25구_2차_공표시점반영(ML용)   위와 같음 + 서울 구 고용 월별 추정
  분기/반기/연간_1차_결측보완      V074 연결, 파생 비율 (기준기간 기준)
  연결계수                연결·역산에 쓴 계수와 금리차

처리 방법 (데이터 시트에는 표시하지 않고 설명서 '보완 내역'에 지역·기간별로 기록)
  원자료              공표값 그대로
  0채움_합계대조       KOSIS가 0인 행을 생략한 것. 합계 - 다른 주택유형 = 0 으로 확인
  차분복원_누계0확인   인허가 누계의 빈 달이 0으로 확인되어 다음 달 월별값을 계산할 수 있게 됨
  0채움_구합계대조     서울 25개 구 합계가 서울 값과 같아 빠진 구는 0 (미분양)
  역산_지수변화        전세가율 = 기준월 전세가율 x 전세지수 변화 / 매매지수 변화
  대용_금리차          은행 주담대금리 + 겹치는 2년 평균 금리차
  전신계열             주택·상가가치전망CSI(2008.07~2012.12)
  연결_구계열          월세가격지수(구)를 2015.06 비율로 이음 (V001 시도)
  역산_권역_학습용     강남·강북 권역 구계열 변화율로 역산. 학습에만 쓰고 평가에 쓰지 않음 (V001 서울 구)
  연결_신계열          가계동향 신계열을 2019년 4개 분기 평균 비율로 이음 (V074)
  추정_상대비          서울 월별 고용 x 최근 공표 반기의 (구 / 서울) 비율 (가용시점 시트)

가용시점 정렬
  행의 월(month) 말일에 알 수 있던 값만 둔다. 월별 변수는 공표 시차(lag)만큼 뒤로 민다.
  분기·반기·연간은 공표월부터 다음 공표 전까지 같은 값을 유지한다(규칙은 설명서 '공표 시차 규칙').
  시차 규칙은 variables.RELEASE / RELEASE_MONTHLY (공식 공표일로 검증, 2026-09-30).
  목표값(Y_*)은 밀지 않는다.

사용법: python impute/build_preprocessed.py
"""

import calendar
import os
import sys
import warnings
from datetime import date

import numpy as np
import pandas as pd

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
sys.path.insert(0, os.path.join(BASE, "merge"))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_panel as bp  # noqa: E402
from check_variables import expected_latest  # noqa: E402
from common import REGIONS, parse_rone_region, raw_path  # noqa: E402
from variables import BY_ID, RELEASE, RELEASE_MONTHLY, drivers  # noqa: E402

warnings.filterwarnings("ignore", category=pd.errors.PerformanceWarning)

START = "2011-01"
OUT = os.path.join(BASE, f"데이터취합_전처리_{date.today():%Y%m%d}.xlsx")

TAG_DOC = {
    "원자료": "공표값 그대로",
    "0채움_합계대조": "KOSIS가 0인 행을 생략해 생긴 빈칸. 같은 달 전체주택 합계 - 다른 유형(단독·다가구·다세대·연립) = 0 으로 확인",
    "차분복원_누계0확인": "인허가 연초누계의 빈 달이 0으로 확인돼, 그 다음 달의 월별값(누계 차분)을 계산할 수 있게 된 셀",
    "0채움_구합계대조": "미분양이 없는 구는 공표하지 않음. 같은 달 공표된 구 합계 = 서울 값으로 확인",
    "역산_지수변화": "전세가율 = 첫 공표월 전세가율 x (전세지수 변화 ÷ 매매지수 변화). 세 계열이 같은 표본",
    "대용_금리차": "은행 주담대금리(V056) + 공표 시작 후 2년 평균 금리차. 대용값",
    "전신계열": "같은 ECOS 표의 전신 '주택·상가가치전망CSI'(2008.07~2012.12). 문항이 달라 D_CSI_전신계열 더미를 함께 둠",
    "연결_구계열": "부동산원 월세가격지수(구)(2012.06=100) x [현행 V001(2015.06) ÷ 구계열(2015.06)]",
    "역산_권역_학습용": "구 V001(2015.06) x 소속 권역(강남/강북) 구계열 변화율. 구별 차이가 없는 값이라 학습에만 사용, 평가 금지",
    "연결_신계열": "가계동향 신계열(2019Q1~) x [구계열 2019년 평균 ÷ 신계열 2019년 평균]",
    "추정_상대비": "서울 월별 값(가용시점) x 최근 공표된 반기의 (구 ÷ 서울 같은 조사월) 비율",
}

LOG = []      # 변수별 처리 내역
PARAMS = []   # 연결계수


def log(vid, col, how, method, filled_s=0, filled_g=0, basis=""):
    LOG.append(dict(ID=vid, 컬럼=col, 처리=how, 방법=method, 채운셀_17시도=filled_s,
                    채운셀_서울구=filled_g, 근거=basis))


def ym(month_int):
    return f"{month_int // 100:04d}-{month_int % 100:02d}"


def set_tag(df, col, mask, label):
    """col 의 꼬리표 컬럼을 만들고(원자료/빈칸) mask 셀에 label 을 단다."""
    t = f"{col}__꼬리표"
    if t not in df.columns:
        df[t] = np.where(df[col].notna(), "원자료", "")
    df.loc[mask, t] = label


def fill(df, col, values, label, where=None):
    """빈칸이면서 values 가 있는 셀만 채운다. 채운 셀 수 반환."""
    values = pd.Series(values, index=df.index)
    mask = df[col].isna() & values.notna()
    if where is not None:
        mask &= where
    if mask.any():
        set_tag(df, col, pd.Series(False, index=df.index), "")  # 꼬리표 컬럼 준비 (원본 기준)
        df.loc[mask, col] = values[mask]
        set_tag(df, col, mask, label)
    return int(mask.sum())


def colof(df, prefix):
    """'V021_' 처럼 시작하는 컬럼 하나"""
    c = [x for x in df.columns if x.startswith(prefix) and "__" not in x]
    return c[0] if c else None


# ============================================================ 1. 0 채움
def housing_zero_verified():
    """주택유형 항등식(합계 = 단독+다가구+다세대+연립+아파트)으로 아파트 빈 달이 0 인 (지역, 연월) 집합."""
    TYPE = {("단독", "단독", "단독"): "단독", ("단독", "다가구", "동수"): "다가구",
            ("다세대", "다세대", "다세대"): "다세대", ("연립", "연립", "연립"): "연립",
            ("아파트", "아파트", "아파트"): "아파트"}
    out, stats = {}, {}
    for tbl in ("DT_MLTM_1948", "DT_MLTM_5387", "DT_MLTM_5373"):
        d = pd.read_csv(raw_path("supplement", f"kosis_주택건설_소분류_{tbl}"), encoding="utf-8-sig", dtype=str)
        d = d[d["C1_NM"].isin(REGIONS)].copy()
        d["v"] = pd.to_numeric(d["DT"], errors="coerce")
        d["t"] = [("합계" if c2.startswith(("합계", "계(")) else TYPE.get((c2, c3, c4)))
                  for c2, c3, c4 in zip(d["C2_NM"], d["C3_NM"], d["C4_NM"])]
        w = d.pivot_table(index=["C1_NM", "PRD_DE"], columns="t", values="v", aggfunc="first")
        parts = ["단독", "다가구", "다세대", "연립", "아파트"]
        full = w.dropna(subset=parts + ["합계"])
        ident = int(((full["합계"] - full[parts].sum(axis=1)).abs() < 0.5).sum())
        miss = w[w["아파트"].isna() & w["합계"].notna()]
        resid = miss["합계"] - miss[["단독", "다가구", "다세대", "연립"]].sum(axis=1, min_count=1).fillna(0)
        zero = resid[resid.abs() < 0.5].index
        out[tbl] = {(r, f"{p[:4]}-{p[4:6]}") for r, p in zero}
        stats[tbl] = (ident, len(full), len(zero), len(miss))
    return out, stats


def zero_fills(sido, gu):
    verified, stats = housing_zero_verified()
    months = sido["month"].map(ym)
    key = list(zip(sido["region"], months))
    names = {"DT_MLTM_1948": "인허가", "DT_MLTM_5387": "착공", "DT_MLTM_5373": "준공"}

    # 착공·준공: 월계라 빈 달 = 0
    for vid, tbl in (("V022", "DT_MLTM_5387"), ("V023", "DT_MLTM_5373")):
        col = colof(sido, vid + "_")
        ok = pd.Series([k in verified[tbl] for k in key], index=sido.index)
        n = fill(sido, col, pd.Series(0.0, index=sido.index), "0채움_합계대조", where=ok)
        ident, full, z, miss = stats[tbl]
        log(vid, col, "0 채움", "빈 달 중 합계 - 다른 유형 = 0 인 달을 0으로", n, 0,
            f"항등식 성립 {ident}/{full}개월, 아파트 빈 달 {miss}개 중 {z}개가 0으로 확인")

    # 인허가: 연초누계를 다시 만들어 빈 누계를 0으로 두고 월별 차분
    col = colof(sido, "V021_")
    raw = pd.read_csv(raw_path("kosis", "DT_MLTM_1948"), encoding="utf-8-sig", dtype=str)
    raw = raw[(raw["C4_NM"] == "아파트") & raw["C1_NM"].isin(REGIONS)]
    cum = {(r, f"{p[:4]}-{p[4:6]}"): float(v) for r, p, v in zip(raw["C1_NM"], raw["PRD_DE"], raw["DT"])
           if v not in ("", "-")}
    for k in verified["DT_MLTM_1948"]:
        cum.setdefault(k, 0.0)
    first = {}
    for (r, m) in cum:
        first[r] = min(first.get(r, m), m)
    new = []
    for r, m in key:
        c = cum.get((r, m))
        if c is None:
            new.append(np.nan)
        elif m.endswith("-01") or m == first.get(r):
            new.append(c)
        else:
            prev = cum.get((r, f"{m[:4]}-{int(m[5:]) - 1:02d}"))
            new.append(np.nan if prev is None else c - prev)
    new = pd.Series(new, index=sido.index)
    before = sido[col].isna()
    zero_mask = before & new.notna() & (new == 0)
    other_mask = before & new.notna() & (new != 0)
    set_tag(sido, col, pd.Series(False, index=sido.index), "")
    sido.loc[zero_mask | other_mask, col] = new[zero_mask | other_mask]
    set_tag(sido, col, zero_mask, "0채움_합계대조")
    set_tag(sido, col, other_mask, "차분복원_누계0확인")
    ident, full, z, miss = stats["DT_MLTM_1948"]
    log("V021", col, "0 채움 + 차분 복원", "빈 누계를 합계 대조로 0 확인 후 월 차분 재계산",
        int(zero_mask.sum() + other_mask.sum()), 0,
        f"항등식 성립 {ident}/{full}개월. 0 {int(zero_mask.sum())}셀, 복원 {int(other_mask.sum())}셀")

    # 미분양 서울 구: 구 합계 = 서울 값인 달에 빠진 구 = 0
    col = colof(gu, "V038_")
    seoul = sido[sido["region"] == "서울"].set_index("month")[colof(sido, "V038_")]
    gsum = gu.groupby("month")[col].sum(min_count=1)
    ok_months = [m for m in gsum.index if pd.notna(seoul.get(m)) and abs(gsum[m] - seoul[m]) < 0.5]
    n = fill(gu, col, pd.Series(0.0, index=gu.index), "0채움_구합계대조", where=gu["month"].isin(ok_months))
    log("V038", col, "0 채움", "구 합계 = 서울 값인 달의 빈 구를 0으로", 0, n,
        f"구 합계 = 서울 값 {len(ok_months)}/{int(seoul.notna().sum())}개월")



# ============================================================ 2. 역산·대용·전신계열
def backcast_jeonse_ratio(df, is_gu):
    jr, m, j = colof(df, "V003_"), colof(df, "V002_"), colof(df, "V026_")
    total = 0
    vals = pd.Series(np.nan, index=df.index)
    for r, g in df.groupby("region"):
        g = g.sort_values("month")
        valid = g[jr].dropna()
        if valid.empty:
            continue
        f = valid.index[0]
        base_jr, base_m, base_j = g.at[f, jr], g.at[f, m], g.at[f, j]
        before = g.index[g["month"] < g.at[f, "month"]]
        vals[before] = base_jr * (g.loc[before, j] / base_j) / (g.loc[before, m] / base_m)
    total = fill(df, jr, vals, "역산_지수변화")
    log("V003", jr, "역산", "첫 공표월 전세가율 x (전세지수 변화 ÷ 매매지수 변화)",
        0 if is_gu else total, total if is_gu else 0, "같은 표본의 세 계열(전세가율·전세지수·매매지수)")


def proxy_rates(df, is_gu):
    base = colof(df, "V056_")
    for vid, w0, w1 in (("V057", 201501, 201612), ("V059", 201401, 201512)):
        c = colof(df, vid + "_")
        nat = df.drop_duplicates("month").set_index("month")
        win = nat[(nat.index >= w0) & (nat.index <= w1)]
        spread = float((win[c] - win[base]).mean())
        n = fill(df, c, df[base] + spread, "대용_금리차")
        if not is_gu:
            PARAMS.append(dict(항목=f"{vid} 금리차", 값=round(spread, 4),
                               설명=f"{ym(w0)}~{ym(w1)} 평균({vid} - V056), %p"))
        log(vid, c, "대용", f"V056 + 금리차 {spread:+.3f}%p ({ym(w0)}~{ym(w1)} 평균)",
            0 if is_gu else n, n if is_gu else 0, "공표 시작 전 구간만. 대용값")


def csi_predecessor(df, is_gu):
    s = pd.read_csv(raw_path("supplement", "ecos_주택상가가치전망CSI"), encoding="utf-8-sig", dtype=str)
    df["D_CSI_전신계열"] = 0
    for vid, code in (("V061", "F0001"), ("V062", "F0002"), ("V063", "F0003")):
        c = colof(df, vid + "_")
        m = s[s["ITEM_CODE2"] == code].set_index("TIME")["DATA_VALUE"].astype(float)
        vals = df["month"].astype(str).map(m)
        n = fill(df, c, vals, "전신계열")
        df.loc[vals.notna() & df[f"{c}__꼬리표"].eq("전신계열"), "D_CSI_전신계열"] = 1
        log(vid, c, "전신계열 연결", "주택·상가가치전망CSI(FMDE) 2008.07~2012.12를 그대로 연결 + 구조변화 더미",
            0 if is_gu else n, n if is_gu else 0, "겹치는 달이 없어 비율 연결 불가. 문항 차이는 D_CSI_전신계열로 흡수")


# ============================================================ 3. 목표변수 V001
def v001_targets(sido, gu):
    old = pd.read_csv(raw_path("supplement", "rone_월세가격지수_구"), encoding="utf-8-sig", dtype=str)
    old["v"] = old["DTA_VAL"].astype(float)
    old["ym"] = old["WRTTIME_IDTFR_ID"].str[:4].astype(int) * 100 + old["WRTTIME_IDTFR_ID"].str[4:].astype(int)
    col = colof(sido, "V001_")

    # 시도: 구계열 비율 연결
    sido["Y_V001_월세통합가격지수(구지수연결)"] = sido[col]
    oldr = {}
    for nm, g in old.groupby("CLS_FULLNM"):
        reg = parse_rone_region(nm)
        if reg and reg[1] == "":
            oldr[reg[0]] = g.set_index("ym")["v"]
    vals = pd.Series(np.nan, index=sido.index)
    for r, s in oldr.items():
        cur = sido[(sido["region"] == r) & (sido["month"] == 201506)][col]
        if cur.empty or pd.isna(cur.iloc[0]) or 201506 not in s.index:
            continue
        k = float(cur.iloc[0]) / s[201506]
        PARAMS.append(dict(항목=f"V001 연결계수 {r}", 값=round(k, 6),
                           설명=f"현행 V001(2015.06)={float(cur.iloc[0]):.3f} ÷ 구계열(2015.06)={s[201506]:.3f}"))
        idx = sido.index[(sido["region"] == r) & (sido["month"] < 201506)]
        vals[idx] = sido.loc[idx, "month"].map(s) * k
    n = fill(sido, "Y_V001_월세통합가격지수(구지수연결)", vals, "연결_구계열")
    log("V001", "Y_V001_월세통합가격지수(구지수연결)", "비율 연결 (목표변수)", "구계열 x [현행(2015.06) ÷ 구계열(2015.06)]", n, 0,
        f"{len(oldr)}개 시도(서울·경기·인천 2011~, 5개 광역시 2012.05~). 나머지 9개 시도는 원자료 없음")

    # 서울 구: 권역 역산 (학습용)
    zone = {}
    pr = pd.read_csv(raw_path("rone", "A_2024_00054"), encoding="utf-8-sig", dtype=str)
    for p in pr["CLS_FULLNM"].dropna().unique():
        parts = p.split(">")
        if parts[0] == "서울" and len(parts) == 4:
            zone[parts[3]] = "서울>" + parts[1]
    colg = colof(gu, "V001_")
    gu["Y_V001_월세통합가격지수(권역역산_학습용)"] = gu[colg]
    vals = pd.Series(np.nan, index=gu.index)
    for g, z in zone.items():
        s = old[old["CLS_FULLNM"] == z].set_index("ym")["v"]
        base = gu[(gu["region"] == g) & (gu["month"] == 201506)][colg]
        if base.empty or pd.isna(base.iloc[0]):
            continue
        idx = gu.index[(gu["region"] == g) & (gu["month"] < 201506)]
        vals[idx] = float(base.iloc[0]) * gu.loc[idx, "month"].map(s) / s[201506]
    n = fill(gu, "Y_V001_월세통합가격지수(권역역산_학습용)", vals, "역산_권역_학습용")
    gu["Y_평가사용가능"] = gu[colg].notna().astype(int)
    sido["Y_평가사용가능"] = sido[col].notna().astype(int)
    for zn in ("서울>강남지역", "서울>강북지역"):
        s = old[old["CLS_FULLNM"] == zn].set_index("ym")["v"]
        PARAMS.append(dict(항목=f"V001 권역 기준값 {zn}", 값=round(s[201506], 4),
                           설명="구계열 2015.06. 구 역산 = 구 V001(2015.06) x 권역(t) ÷ 이 값"))
    log("V001", "Y_V001_월세통합가격지수(권역역산_학습용)", "권역 역산 (목표변수, 학습 전용)",
        "구 V001(2015.06) x 소속 권역 구계열(t) ÷ 권역 구계열(2015.06)", 0, n,
        "강남 11개·강북 14개 구. 2015.06 이전 구별 차이가 없는 값이라 평가에서 제외(Y_평가사용가능=0)")


# ============================================================ 4. 분기·반기·연간
def lowfreq(sido_base, gu_base):
    q, _ = bp.lowfreq_panel("Q")
    y, _ = bp.lowfreq_panel("Y")
    h, _ = bp.lowfreq_panel("H")
    colvid = {}
    for df in (q, y, h):
        for c in df.columns[3:]:
            colvid[c] = "V013H" if c.startswith("V013") else c[:4]

    # V074 신계열 연결
    old_c = colof(q, "V074_")
    new = pd.read_csv(raw_path("supplement", "kosis_가계동향_신계열"), encoding="utf-8-sig", dtype=str)
    new = new.set_index(new["PRD_DE"].str[:4] + "Q" + new["PRD_DE"].str[4:6].astype(int).astype(str))["DT"].astype(float)
    qn = q.set_index("기간")
    ov = [f"2019Q{i}" for i in range(1, 5)]
    k = float(qn.loc[ov, old_c].mean() / new[ov].mean())
    PARAMS.append(dict(항목="V074 연결계수", 값=round(k, 6), 설명="구계열 2019년 4개 분기 평균 ÷ 신계열 같은 기간 평균"))
    q["V074_연결_월평균소득_원"] = q[old_c]
    ext = pd.Series(np.nan, index=q.index)
    later = q["기간"] > "2019Q4"
    ext[later] = q.loc[later, "기간"].map(new) * k
    n = fill(q, "V074_연결_월평균소득_원", ext, "연결_신계열")
    colvid["V074_연결_월평균소득_원"] = "V074"
    log("V074", "V074_연결_월평균소득_원", "비율 연결", f"2020Q1~ = 신계열 x {k:.4f}", 0, 0,
        f"2019년 4개 분기 겹침. {n}개 분기 연장 (분기 시트)")

    # 파생 비율 (기준기간 기준)
    y["_연월"] = y["기간"].astype(str) + "-12"
    hh_s = sido_base.assign(_연월=sido_base["month"].map(ym)).set_index(["region", "_연월"])[colof(sido_base, "V010_")]
    hh_g = gu_base.assign(_연월=gu_base["month"].map(ym)).set_index(["region", "_연월"])[colof(gu_base, "V010_")]
    hh = pd.Series([hh_g.get((g, m)) if g else hh_s.get((s, m)) for s, g, m in zip(y["시도"], y["구"], y["_연월"])],
                   index=y.index, dtype=float)
    y["V072_주택수÷가구수"] = y[colof(y, "V072_총조사_주택수")] / y[colof(y, "V072_총조사_가구수")]
    colvid["V072_주택수÷가구수"] = "V072"
    derived = ["V072_주택수÷가구수"]
    for c in [x for x in y.columns if x.startswith(("V024_주거용신축허가_동수", "V076_", "V077_")) and "__" not in x]:
        nc = c.rsplit("_", 1)[0] + "_천세대당"
        y[nc] = y[c] / hh * 1000
        colvid[nc] = c[:4]
        derived.append(nc)
    nat = pd.read_csv(raw_path("kosis", "DT_1C96"), encoding="utf-8-sig", dtype=str)
    nat = nat[nat["C1"] == "00"]
    for vid, itm, c in (("V015", "T1", colof(y, "V015_")), ("V017", "T3", colof(y, "V017_"))):
        s = nat[nat["ITM_ID"] == itm].set_index("PRD_DE")["DT"].astype(float)
        nc = f"{vid}_전국대비"
        y[nc] = y[c] / y["기간"].astype(str).map(s)
        colvid[nc] = vid
        derived.append(nc)
    y = y.drop(columns="_연월")
    log("파생", ", ".join(derived), "파생 비율",
        "주택수÷가구수 / 12월 세대수(V010) 1천 세대당 / 전국 대비(전국=1)", 0, 0, "기준연도 기준으로 계산 후 공표월부터 반영")
    return q, y, h, colvid


# ============================================================ 5. 가용시점 정렬
def month_end(m):
    y_, mo = m // 100, m % 100
    return date(y_, mo, calendar.monthrange(y_, mo)[1])


def align(df, colvid, keep):
    """월별 변수를 공표 시차만큼 지역별로 민다. keep 컬럼(식별자·목표값)은 그대로."""
    out = df[keep].copy()
    df = df.sort_values(["region", "month"])
    for c in df.columns:
        if c in keep:
            continue
        base_c = c.split("__")[0]
        vid = colvid.get(base_c)
        lag = RELEASE_MONTHLY[vid]["lag"] if vid in RELEASE_MONTHLY else 0
        out[c] = df.groupby("region")[c].shift(lag) if lag else df[c]
    return out


def stepfill(target, low, cols, colvid, region_field):
    """분기·반기·연간 값을 '행의 월 말일에 공표돼 있던 최신 기'로 채우고 기준기간 꼬리표를 단다."""
    months = sorted(target["month"].unique())
    for c in cols:
        vid = colvid[c]
        rule = RELEASE[vid]
        sample = str(low["기간"].dropna().iloc[0])
        per = {m: expected_latest(sample, rule, month_end(m)) for m in months}
        pmap = target["month"].map(per)
        if region_field is None:
            s = low.drop_duplicates("기간").set_index("기간")[c]
            vals = pmap.map(s)
        else:
            s = low.set_index([region_field, "기간"])[c]
            vals = pd.Series([s.get((r, p)) for r, p in zip(target["region"], pmap)], index=target.index, dtype=float)
        target[c] = vals
        target[f"{c}__기준기간"] = np.where(vals.notna(), pmap, "")
        tag_src = f"{c}__꼬리표"
        if tag_src in low.columns:
            tags = low.set_index("기간")[tag_src] if region_field is None else None
            if tags is not None:
                target[tag_src] = np.where(vals.notna(), pmap.map(tags.groupby(level=0).first()), "")


def seoul_monthly_estimate(gu_al, sido_al, h, colvid):
    """V013 서울 구 월별 추정 = 서울 월별(가용시점) x 최근 공표 반기의 (구 ÷ 서울 조사월) 비율"""
    s_emp = sido_al[sido_al["region"] == "서울"].set_index("month")
    raw_s = {}
    for item, prefix in (("고용률", "V013_고용률"), ("취업자", "V013_취업자")):
        raw_s[item] = colof(s_emp, prefix)
    # 서울 월별 원값(기준기간)은 정렬 전 값이 필요 -> sido_al 은 1개월 밀려 있으므로 되돌려 쓴다
    for item, hcol in (("고용률", colof(h, "V013_서울구반기_고용률")), ("취업자", colof(h, "V013_서울구반기_취업자"))):
        s_now = s_emp[raw_s[item]]                     # 가용시점 서울 월별
        s_ref = s_now.shift(-RELEASE_MONTHLY["V013"]["lag"])  # 기준월 서울 월별
        ratio = {}
        for _, r in h.iterrows():
            ref_m = int(r["기간"][:4]) * 100 + (4 if r["기간"].endswith("H1") else 10)
            if pd.notna(r[hcol]) and pd.notna(s_ref.get(ref_m)) and s_ref.get(ref_m):
                ratio[(r["구"], r["기간"])] = r[hcol] / s_ref[ref_m]
        per = {m: expected_latest("2021H1", RELEASE["V013H"], month_end(m)) for m in gu_al["month"].unique()}
        nc = f"V013_서울구_{item}_월별추정"
        vals = [s_now.get(m) * ratio[(g, per[m])] if (g, per[m]) in ratio and pd.notna(s_now.get(m)) else np.nan
                for g, m in zip(gu_al["region"], gu_al["month"])]
        gu_al[nc] = vals
        gu_al[f"{nc}__꼬리표"] = np.where(pd.notna(vals), "추정_상대비", "")
        log("V013", nc, "월별 추정 (가용시점 시트)", "서울 월별(가용시점) x 최근 공표 반기 (구 ÷ 서울 조사월)",
            0, int(pd.notna(vals).sum()), "반기 계단값(V013_서울구반기_*)도 함께 둠")


# ============================================================ 실행
def reorder(df):
    """꼬리표·기준기간 컬럼을 값 컬럼 바로 오른쪽으로"""
    base = [c for c in df.columns if "__" not in c]
    out = []
    for c in base:
        out.append(c)
        out += [x for x in df.columns if x.startswith(c + "__")]
    out += [c for c in df.columns if c not in out]
    return df[out]


def main():
    sido_f, gu_f, nation, meta = bp.monthly_frames()
    colvid = {m["컬럼"]: m["id"] for m in meta}
    allm = set()
    for s in list(sido_f.values()) + list(gu_f.values()):
        allm |= set(s.index.get_level_values("연월"))
    months = [f"{p.year:04d}-{p.month:02d}" for p in pd.period_range(START, max(allm), freq="M")]
    sido = bp.build_monthly(sido_f, REGIONS, "광역지자체", "17개 시도",
                            lambda r: f"{bp.SIDO_CODE[r]}000000", months, nation)
    gu = bp.build_monthly(gu_f, bp.GU_ORDER, "자치구", "서울 25개 구",
                          lambda r: f"{bp.GU_CODE[r]}000", months, nation)
    print(f"[기반] 17시도 {sido.shape}, 서울구 {gu.shape}")

    zero_fills(sido, gu)
    for df, is_gu in ((sido, False), (gu, True)):
        backcast_jeonse_ratio(df, is_gu)
        proxy_rates(df, is_gu)
        csi_predecessor(df, is_gu)
    v001_targets(sido, gu)

    # 보완 불가 (빈칸 유지) - 데이터 시트에는 표시 열을 두지 않고 설명서에만 기록
    log("V066", colof(sido, "V066_"), "보완 안 함",
        "2011~2019년 빈칸 유지: 월별 전체근로자 임금 표 없음(KOSIS DT_118N_MON048은 연간·상용근로자만)", 0, 0, "")
    log("V071", colof(sido, "V071_"), "보완 안 함", "2015.06 이전 빈칸 유지: 통계 작성 전", 0, 0, "")

    # 서울 상위지역 변수 (구 패널)
    up = ("V005", "V013", "V021", "V022", "V023", "V031", "V032", "V037", "V050")
    seoul = sido[sido["region"] == "서울"].set_index("month")
    added = []
    for c in [x for x in sido.columns if x[:4] in up and "__" not in x]:
        nc = f"서울_{c}"
        gu[nc] = gu["month"].map(seoul[c])
        if f"{c}__꼬리표" in seoul.columns:
            gu[f"{nc}__꼬리표"] = gu["month"].map(seoul[f"{c}__꼬리표"])
        colvid[nc] = c[:4]
        added.append(nc)
    log("상위지역", f"서울_* {len(added)}개", "상위지역 변수", "구 단위가 없는 시도 변수에 서울 값을 붙임(이름 앞 '서울_')",
        0, 0, ", ".join(sorted({c[3:7] for c in added})))

    colvid["D_CSI_전신계열"] = "V061"
    q, y, h, lcol = lowfreq(sido, gu)

    # ---------------- 가용시점 정렬
    keep = ["panel", "region", "region_code", "month"] + [c for c in sido.columns if c.startswith("Y_")]
    sido_al = align(sido, colvid, [c for c in keep if c in sido.columns])
    keep_g = ["panel", "region", "region_code", "month"] + [c for c in gu.columns if c.startswith("Y_")]
    gu_al = align(gu, colvid, keep_g)
    qcols = [c for c in q.columns[3:] if "__" not in c]
    stepfill(sido_al, q, qcols, lcol, None)
    stepfill(gu_al, q, qcols, lcol, None)
    ys, yg = y[y["구"] == ""], y[y["구"] != ""]
    ycols_s = [c for c in y.columns[3:] if "__" not in c and ys[c].notna().any()]
    ycols_g = [c for c in y.columns[3:] if "__" not in c and yg[c].notna().any()]
    stepfill(sido_al, ys.rename(columns={}), ycols_s, lcol, "시도")
    stepfill(gu_al, yg, ycols_g, lcol, "구")
    hcols = [c for c in h.columns[3:] if "__" not in c]
    stepfill(gu_al, h, hcols, lcol, "구")
    seoul_monthly_estimate(gu_al, sido_al, h, lcol)
    write(sido, gu, sido_al, gu_al, q, y, h, months, colvid, lcol)


FREQ_KO = {"M": "월", "Q": "분기", "H": "반기", "Y": "연"}
SHEETS = {"s1": "17시도_1차_결측보완", "g1": "서울25구_1차_결측보완",
          "s2": "17시도_2차_공표시점반영(ML용)", "g2": "서울25구_2차_공표시점반영(ML용)",
          "q": "분기_1차_결측보완", "h": "반기_1차_결측보완", "y": "연간_1차_결측보완"}
ID_COLS = {"panel": "패널 구분", "region": "지역(시도명 또는 구명)", "region_code": "행정표준코드",
           "month": "연월(YYYYMM)", "시도": "지역(시도)", "구": "지역(서울 구, 시도 행은 빈칸)", "기간": "기준기간"}


def period_key(p):
    p = str(p)
    if len(p) == 6 and p.isdigit():
        return int(p[:4]) * 12 + int(p[4:])
    return int(p[:4]) * 12 + (int(p[-1]) * (3 if "Q" in p else 6) if ("Q" in p or "H" in p) else 12)


def fmt_period(p):
    p = str(p)
    return f"{p[:4]}.{p[4:]}" if len(p) == 6 and p.isdigit() else p


def ranges(periods):
    """연속 구간으로 묶어 '2011.01~2012.12' 형태 문자열 (월 1, 분기 3, 반기 6, 연 12 간격이면 연속)"""
    ps = sorted(set(periods), key=period_key)
    p0 = str(ps[0])
    step = 1 if (len(p0) == 6 and p0.isdigit()) else 3 if "Q" in p0 else 6 if "H" in p0 else 12
    out, start, prev = [], ps[0], ps[0]
    for x in ps[1:]:
        if period_key(x) - period_key(prev) != step:
            out.append((start, prev))
            start = x
        prev = x
    out.append((start, prev))
    return ", ".join(fmt_period(a) if a == b else f"{fmt_period(a)}~{fmt_period(b)}" for a, b in out)


def fill_details(frames):
    """꼬리표 열에서 '어느 지역·기간을 어떤 방법으로 채웠나'를 요약한 표"""
    rows = []
    for key, df, pcol, rcol, n_regions in frames:
        for t in [c for c in df.columns if c.endswith("__꼬리표")]:
            col = t[:-len("__꼬리표")]
            if col.startswith("서울_"):
                continue  # 서울 시도 값의 복사본. 원 변수 쪽에 기록됨
            if key == "g2" and "월별추정" not in col:
                continue  # 2차 시트의 나머지 값은 1차를 민 것이라 1차 쪽에 기록됨
            sub = df[~df[t].isin(["원자료", ""]) & df[t].notna()]
            for label, g in sub.groupby(t):
                by_rng = {}
                for reg, gg in g.groupby(rcol) if rcol else [("전국", g)]:
                    by_rng.setdefault(ranges(gg[pcol].tolist()), []).append(str(reg) if reg != "" else "전국")
                for rng, regs in by_rng.items():
                    who = f"전 지역({len(regs)}개)" if n_regions and len(regs) == n_regions else ", ".join(regs)
                    rows.append({"시트": SHEETS[key], "컬럼": col, "처리 방법": label, "대상 지역": who,
                                 "보완 기간": rng, "셀 수": int(len(g[g[rcol].isin(regs)]) if rcol else len(g))})
    return pd.DataFrame(rows)


def lag_text(vid, freq):
    if vid in RELEASE_MONTHLY and freq == "M":
        L = RELEASE_MONTHLY[vid]["lag"]
        return "t월 행 = t월 값 (공표 시차 없음)" if L == 0 else f"t월 행 = t-{L}월 값 ({L}개월 밀림)"
    if vid in RELEASE:
        L, unit = RELEASE[vid]["lag"], {"Q": "분기", "H": "반기", "Y": "연도"}.get(freq, "기간")
        if L < 0:
            return f"해당 {unit} 첫 달부터 그 값 사용(전망 조사), 다음 공표 전까지 유지"
        return f"{unit} 종료 후 {L}개월째 달부터 그 {unit} 값 사용, 다음 공표 전까지 유지 (계단식)"
    return "-"


def data_catalog(frames_cols, colvid, lcol, details):
    """설명서의 '데이터 목록': 데이터 시트의 모든 열 한 줄씩"""
    method = {}
    for e in LOG:
        for c in [x.strip() for x in str(e["컬럼"]).split(",")]:
            method.setdefault(c, []).append(f"{e['처리']}: {e['방법']}")
    filled = details.groupby("컬럼")["셀 수"].sum().to_dict() if len(details) else {}
    order, where = [], {}
    for key, cols in frames_cols:
        for c in cols:
            if c not in where:
                order.append(c)
            where.setdefault(c, []).append(SHEETS[key])
    rows = []
    for c in order:
        sheets = " / ".join(where[c])
        if c in ID_COLS:
            rows.append({"컬럼명": c, "Master ID": "-", "변수명(데이터사전)": ID_COLS[c], "주기": "-",
                         "수록 시트": sheets, "지역 단위": "-", "출처(기관·표)": "-", "전처리 방법": "식별 열",
                         "보완 셀 수": "", "공표 시기 (공식 근거)": "-", "2차 시트 시차 규칙": "-", "사전 대비": "-"})
            continue
        base = c[3:] if c.startswith("서울_") else c
        vid = colvid.get(c) or colvid.get(base) or lcol.get(c) or lcol.get(base)
        if c.startswith("V013_서울구_") and "월별추정" in c:
            vid = "V013H"
        if c.startswith("Y_"):
            vid = "V001"
        if c == "D_CSI_전신계열":
            vid = "V061"
        v = BY_ID.get(vid, {})
        freq = "M" if (c in colvid or base in colvid) and not c.startswith("V013_서울구반기") else v.get("freq", "")
        if c.startswith("V013_서울구반기"):
            vid, freq = "V013H", "H"
        rule = RELEASE_MONTHLY.get(vid) if freq == "M" else RELEASE.get(vid)
        how = method.get(c) or method.get(base) or []
        if c.startswith("서울_"):
            how = ["상위지역 변수: 서울 시도 값을 25개 구 모두에 동일하게 붙임 (구 단위 통계 없음)"] + how
        if c == "D_CSI_전신계열":
            how = ["구조변화 더미: V061~V063을 전신계열(주택·상가가치전망CSI)로 채운 달(2011.01~2012.12) = 1, 그 외 0"]
        if c == "Y_평가사용가능":
            how = ["평가 표시: 공식 V001이 있는 행 = 1. 성능 평가는 1인 행만 사용"]
        if c.startswith("V013_서울구_") and "월별추정" in c:
            how = ["2차 시트 전용 추정: 서울 월별 고용(2차 기준) x 최근 공표된 반기의 (구 ÷ 서울 조사월) 비율"]
        if not how:
            how = ["보완 없음 (원자료 그대로)"]
        if c.startswith("V013_서울구_") and "월별추정" in c:
            lag = "서울 월별 고용은 t-1월 값, 반기 비율은 t월 말까지 공표된 최근 반기"
        elif c.startswith("Y_"):
            lag = "밀지 않음 (목표값)"
        elif c in ("Y_평가사용가능",):
            lag = "밀지 않음"
        else:
            lag = lag_text(vid, freq)
        rows.append({
            "컬럼명": c, "Master ID": vid or "-", "변수명(데이터사전)": v.get("name", "-"),
            "주기": FREQ_KO.get(freq, freq), "수록 시트": sheets,
            "지역 단위": "서울 시도 값(상위지역)" if c.startswith("서울_") else v.get("level", "-"),
            "출처(기관·표)": f"{v.get('source', '')} {v.get('table', '')}".strip() or "-",
            "전처리 방법": " / ".join(dict.fromkeys(how)),
            "보완 셀 수": filled.get(c, filled.get(base, "")),
            "공표 시기 (공식 근거)": rule["evidence"] if rule else "-",
            "2차 시트 시차 규칙": lag,
            "사전 대비": f"{rule['verdict']} (사전: {rule['dict']})" if rule else "-",
        })
    df = pd.DataFrame(rows)
    df.insert(0, "번호", range(1, len(df) + 1))
    return df


# 2차(ML용) 시트에서 빼는 열 - 다른 열로 완전히 계산되거나, 상관이 매우 높거나, 쓸 수 없는 열.
# 1차 시트에는 모두 남긴다. 데이터사전 변수는 하나도 빠지지 않는다(변수마다 최소 1개 열 유지).
ML_EXCLUDE = {
    "V006_전월세_전세_건": "전체 = 전세 + 월세 가 100% 성립",
    "V011_20-39세_명": "20-39세 = 20-29세 + 30-39세 가 100% 성립",
    "V035_아파트거래_면적": "동·호수와 상관 0.97",
    "V036_아파트매매거래_면적": "동·호수와 상관 0.97",
    "V076_공공+민간총계(DT_MLTM_7174)": "2024년 1년치뿐이라 학습 불가",
    "V076_공공임대계(DT_MLTM_7174)": "2024년 1년치뿐이라 학습 불가",
    "V076_민간임대계(DT_MLTM_7174)": "2024년 1년치뿐이라 학습 불가",
    "V074_가계동향_도시2인이상_월평균소득_원": "2019Q4 종료. V074_연결 사용",
    "V069_예금취급기관_가계대출": "주택관련대출과 수준 상관 0.98. 주택관련대출 사용",
    "V072_총조사_가구수": "지역 규모만 반영(주택수와 상관 0.92). V072_주택수÷가구수 사용",
    "V072_총조사_주택수": "지역 규모만 반영. V072_주택수÷가구수 사용",
    "V024_주거용신축허가_동수_동": "지역 규모에 끌려감. 천세대당 사용",
    "V076_민간임대(DT_MLTM_6827합계)_호": "지역 규모에 끌려감. 천세대당 사용",
    "V077_임대주택건설공급_총계_호": "지역 규모에 끌려감. 천세대당 사용",
    "V015_1인당GRDP_천원": "전국 공통 추세 포함. 전국대비 사용",
    "V017_1인당가계총처분가능소득_천원": "전국 공통 추세 포함. 전국대비 사용",
    "V013_취업자": "인구 규모 반영(V009와 역할 중복). 고용률 사용",
    "V013_서울구반기_": "월별 추정 고용률과 같은 정보. V013_서울구_고용률_월별추정 사용",
    "V013_서울구_취업자_월별추정": "고용률과 지역 내 상관 0.83. 고용률 사용",
}


def ml_excluded(col):
    """2차(ML) 시트 제외 사유. 서울_ 상위지역 복사본도 같은 규칙. 제외 아니면 None"""
    base = col[3:] if col.startswith("서울_") else col
    for k, why in ML_EXCLUDE.items():
        if base.startswith(k):
            return why
    return None


MANUAL = ["1_개요", "2_데이터목록", "3_전처리방법", "4_공표시차규칙"]


def write(sido, gu, sido_al, gu_al, q, y, h, months, colvid, lcol):
    details = fill_details([
        ("s1", sido, "month", "region", 17), ("g1", gu, "month", "region", 25),
        ("g2", gu_al, "month", "region", 25), ("q", q, "기간", None, None)])
    # V001 공식 원본 열은 목표값(Y_V001_월세통합가격지수…)에 들어 있으므로 데이터 시트에서 뺀다(설명변수 아님)
    def clean(df):
        cols = [c for c in df.columns if "__" not in c and not c.startswith("V001_")]
        ids = [c for c in cols if c in ("panel", "region", "region_code", "month")]
        ys = [c for c in cols if c.startswith("Y_")]  # 목표값은 식별 열 바로 뒤
        return df[ids + ys + [c for c in cols if c not in ids and c not in ys]]
    ml = lambda df: df[[c for c in df.columns if not ml_excluded(c)]]
    out = {"s1": clean(sido), "g1": clean(gu), "s2": ml(clean(sido_al)), "g2": ml(clean(gu_al)),
           "q": clean(q), "h": clean(h), "y": clean(y)}
    catalog = data_catalog([(k, list(d.columns)) for k, d in out.items()], colvid, lcol, details)
    catalog.insert(catalog.columns.get_loc("수록 시트") + 1, "ML(2차) 사용",
                   catalog["컬럼명"].map(lambda c: "-" if c in ID_COLS else ("X: " + ml_excluded(c)) if ml_excluded(c) else "O"))
    catalog = catalog.drop(columns=["사전 대비"])
    drv = drivers()
    i = catalog.columns.get_loc("변수명(데이터사전)")
    catalog.insert(i, "동인", catalog["Master ID"].map(lambda x: drv.get(x, ("", ""))[0]))

    # 1_개요
    intro = pd.DataFrame([
        ("생성", f"{date.today()} | 기간 {months[0]} ~ {months[-1]} | 17개 시도 {len(sido):,}행, 서울 25개 구 {len(gu):,}행"),
        ("입력", "processed/ (로데이터 취합본과 같은 입력) + raw/7_보완용_사전외/ (보완용 원자료, 데이터사전 변수 아님)"),
        ("변수 범위", "docs/260928_데이터사전_취합_수정.xlsx 의 Variable_Master 52개 변수. 컬럼명 앞 V### 가 Master ID"),
        ("처리 순서", "로데이터 취합본(빈칸 그대로) → 1차_결측보완 → 2차_공표시점반영(ML용)"),
        ("1차_결측보완", "빈칸만 채움. 원본 값은 바꾸지 않음. 각 행 = 그 달(기준기간)의 값. 어떤 셀을 어떻게 채웠는지는 '3_전처리방법' 시트"),
        ("2차_공표시점반영", "1차 결과를 변수별 공표 시차만큼 밀어, 각 행 = 그 달 말에 실제로 알 수 있던 값. "
                         "분기·반기·연간은 공표월부터 다음 공표 전까지 같은 값. 머신러닝 입력은 이 시트. 규칙은 '4_공표시차규칙' 시트"),
        ("ML(2차) 열 선택", "2차 시트에는 다른 열로 완전히 계산되거나(예: 전체=전세+월세) 상관이 매우 높거나 1년치뿐인 열을 뺐다. "
                         "1차 시트에는 모두 남아 있고, 뺀 열과 사유는 2_데이터목록의 'ML(2차) 사용' 열. 데이터사전 변수는 모두 1개 이상 열로 남음"),
        ("분기·반기·연간", "분기_/반기_/연간_1차_결측보완 시트는 기준기간 기준 원표. 공표 시차를 반영한 2차 값은 따로 시트를 두지 않고 "
                      "17시도·서울25구_2차_공표시점반영(ML용) 시트 안에 열로 들어 있음(공표월부터 다음 공표 전까지 같은 값). "
                      "예: 가계신용 2026년 2분기 값은 8/19 공표 → 2차 시트 2026.08~ 행에 들어감"),
        ("목표변수", "Y_ 로 시작. 2차 시트에서도 밀지 않음. 시도 = Y_V001_월세통합가격지수(구지수연결)(2015.06 이전을 월세가격지수(구)로 비율 연결), "
                 "서울 구 = Y_V001_월세통합가격지수(권역역산_학습용)(2015.06 이전을 강남/강북 권역 구계열로 역산)"),
        ("V001 월세통합지수", "목표변수로만 둔다(설명변수 열 없음). 2015.06 이후는 공식 지수 그대로, 이전은 위 방법으로 채운 값. "
                         "공식 값인지 채운 값인지는 Y_평가사용가능 으로 구분. 지수의 과거값을 입력으로 쓰려면 모델에서 목표값을 직접 시차 처리"),
        ("Y_평가사용가능", "그 행의 목표값이 공식 월세통합지수이면 1(2015.06 이후), 연결·역산으로 채운 값이거나 빈칸이면 0. "
                        "학습은 채운 값까지 전부 쓰고, 성능 평가는 1인 행만 쓴다. 채운 값을 정답으로 평가하면 실제 시장이 아니라 "
                        "보완 방식을 얼마나 따라 했는지를 재게 되기 때문. 특히 서울 구의 역산값은 같은 권역 구들이 똑같이 움직여 "
                        "구별 차이를 평가할 수 없음. 시도 연결값도 기준이 다른 지수를 이은 것이라 보수적으로 0"),
        ("보완 우선순위", "0이 확실한 것 → 같은 표본으로 역산 → 겹치는 시점 비율로 연결 → 대용 → 불가하면 빈칸 유지"),
        ("빈칸", "보완 후에도 남은 빈칸은 0이 아님(미공표·미작성·지역 없음). 세종은 2012.07 출범 전 값이 없음"),
        ("시트 구성", "1_개요 / 2_데이터목록 / 3_전처리방법 / 4_공표시차규칙 / 17시도·서울25구_1차_결측보완 / "
                  "17시도·서울25구_2차_공표시점반영(ML용) / 분기·반기·연간_1차_결측보완 / 연결계수"),
    ], columns=["항목", "내용"])

    # 3_전처리방법: 위 = 유형별 설명, 아래 = 데이터별 내용
    methods = pd.DataFrame([dict(처리방법=k, 설명=v) for k, v in TAG_DOC.items() if k != "원자료"])
    logd = {}
    for e in LOG:
        for c in [x.strip() for x in str(e["컬럼"]).split(",")]:
            logd.setdefault(c, (e["방법"], e["근거"]))
    per_data = details.copy()
    per_data.insert(3, "방법 상세", per_data["컬럼"].map(lambda c: logd.get(c, ("", ""))[0]))
    per_data["근거"] = per_data["컬럼"].map(lambda c: logd.get(c, ("", ""))[1])
    # 보완하지 않은 변수(빈칸 유지)와 시트 전체 처리도 함께 적는다
    extra = []
    for e in LOG:
        if e["처리"].startswith(("보완 안 함", "상위지역", "파생", "비율 연결")) and e["컬럼"] not in set(details["컬럼"]):
            extra.append({"시트": "-", "컬럼": e["컬럼"], "처리 방법": e["처리"], "방법 상세": e["방법"],
                          "대상 지역": "-", "보완 기간": "-", "셀 수": "", "근거": e["근거"]})
    per_data = pd.concat([per_data, pd.DataFrame(extra)], ignore_index=True)
    per_data = per_data[["시트", "컬럼", "처리 방법", "방법 상세", "대상 지역", "보완 기간", "셀 수", "근거"]]

    # 4_공표시차규칙
    rules = []
    for vid, r in {**RELEASE_MONTHLY, **RELEASE}.items():
        monthly = vid in RELEASE_MONTHLY
        v = BY_ID.get(vid, {})
        rules.append({"Master ID": vid, "동인": drv.get(vid, ("", ""))[0],
                      "변수명": v.get("name", ""), "주기": "월" if monthly else FREQ_KO.get(v.get("freq"), ""),
                      "공표 시차(개월)": r["lag"],
                      "공표일": ("말일" if r.get("day") == 31 else f"{r['day']}일경") if r.get("day") else "월 단위만 공식",
                      "2차 시트 적용": "밀지 않음 (목표변수)" if vid == "V001" else lag_text(vid, "M" if monthly else v.get("freq", "")),
                      "공식 근거": r["evidence"], "출처": r["source"]})
    rules = pd.DataFrame(rules)

    # 열 제목을 "ID_무슨 통계_세부_단위" 로 바꾼다 (데이터 시트 + 설명서 안의 열 이름 언급까지 한 번에)
    import re
    from display_names import DISPLAY, disp
    out = {k: df.rename(columns=disp) for k, df in out.items()}
    names = sorted({c for c in DISPLAY if DISPLAY[c] != c}, key=len, reverse=True)
    pat = re.compile("|".join(re.escape(n) for n in names))
    fix = lambda v: pat.sub(lambda m: DISPLAY[m.group(0)], v) if isinstance(v, str) else v  # noqa: E731
    intro, catalog, per_data, methods, rules = [df.apply(lambda s: s.map(fix)) for df in
                                                (intro, catalog, per_data, methods, rules)]

    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        intro.to_excel(xw, sheet_name="1_개요", index=False)
        catalog.to_excel(xw, sheet_name="2_데이터목록", index=False)
        pd.DataFrame({"■ 전처리 방법 (유형별)": []}).to_excel(xw, sheet_name="3_전처리방법", index=False)
        methods.to_excel(xw, sheet_name="3_전처리방법", startrow=1, index=False)
        r2 = len(methods) + 4
        pd.DataFrame({"■ 데이터별 전처리 내용 (어느 지역·기간을 어떤 방법으로)": []}).to_excel(
            xw, sheet_name="3_전처리방법", startrow=r2, index=False)
        per_data.to_excel(xw, sheet_name="3_전처리방법", startrow=r2 + 1, index=False)
        rules.to_excel(xw, sheet_name="4_공표시차규칙", index=False)
        for k, df in out.items():
            df.to_excel(xw, sheet_name=SHEETS[k], index=False)
        pd.DataFrame(PARAMS).to_excel(xw, sheet_name="연결계수", index=False)
        style(xw.book)
    print(f"완료: {OUT}")
    print(f"  데이터 목록 {len(catalog)}열, 데이터별 전처리 {len(per_data)}줄, 공표시차 규칙 {len(rules)}개")


def style(wb):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    head = PatternFill("solid", fgColor="DDE6F0")
    widths = {"1_개요": [18, 120],
              "2_데이터목록": [6, 34, 10, 24, 30, 6, 34, 30, 14, 26, 60, 9, 60, 40],
              "3_전처리방법": [24, 34, 20, 50, 20, 60, 8, 50],
              "4_공표시차규칙": [10, 24, 32, 8, 10, 14, 44, 70, 40]}
    for ws in wb.worksheets:
        if ws.title in MANUAL:
            for i, w in enumerate(widths[ws.title], 1):
                ws.column_dimensions[get_column_letter(i)].width = w
            prev_title = ws.title != "3_전처리방법"  # 3번 시트만 소제목 아래가 헤더
            for n, row in enumerate(ws.iter_rows(), 1):
                first = row[0].value
                is_title = isinstance(first, str) and first.startswith("■")
                is_head = (n == 1 and ws.title != "3_전처리방법") or prev_title
                for c in row:
                    c.font = Font(name="Arial", size=12 if is_title else 10, bold=is_title or is_head)
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    if is_head and c.value is not None:
                        c.fill = head
                prev_title = is_title
            ws.freeze_panes = "A2" if ws.title != "3_전처리방법" else None
            continue
        for i, c in enumerate(ws[1], 1):
            c.font = Font(name="Arial", size=10, bold=True)
            c.fill = head
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(c.value)) * 1.2, 9), 26)
        ws.row_dimensions[1].height = 48
        ws.freeze_panes = "E2" if ws.title.startswith(("17시도", "서울25구")) else "D2"
        ws.auto_filter.ref = ws.dimensions


if __name__ == "__main__":
    main()
