# -*- coding: utf-8 -*-
"""
모형 입력표(xlsx) 작성 — 결정 로그(docs/전처리결정로그_10차.csv)의 '결정' 행을 raw 원자료에 적용한다. variable_card.py 가 부른다.

파일 하나: analysis/output/모형입력표_10차.xlsx
  '타깃'          17개 시도 주 패널. A region, B 결정월, C R_last(원값: 결정월에 아는 마지막 지수 R(t−1)), D~J 는 C열을 참조하는 수식
                  (past1·3·6, G3, r3, G6, r6), K 구간(버퍼/분석)
  '설명변수'       17개 시도. A region, B 결정월, C 구간, 그 뒤 변수마다 '원값(달력월)' 열 + 결정된 변환의 수식 열. 보조 열(세대수·인구)은 필요할 때 자동 추가
  '타깃_서울구'    서울 25개 구 보조 패널(공식 구 지수 2015-06~). 구조 같음. r_h 는 25개 구 평균 대비
  '설명변수_서울구' 구 단위 raw 가 있으면 구 값, 없으면 서울 시도 값을 복제(열 이름에 표시)
  '설명'          열마다 수식의 뜻
행: 지역 × 결정월(2014-01 버퍼 ~ 자료 끝), 지역 블록 안에서 월 순서. 수식은 같은 지역 블록 안의 위·아래 행만 참조하며,
블록 경계를 넘거나 값이 없으면 "" 를 돌려준다. 2016-01 이전 행은 12개월 변화 계산용 버퍼(구간='버퍼').
"""

import os
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEET_PATH = os.path.join(BASE, "analysis", "output", "모형입력표_10차.xlsx")
BUFFER_START = pd.Period("2014-01", "M")
ANALYSIS_START = pd.Period("2016-01", "M")
GREY = PatternFill("solid", fgColor="EEEEEE")
BOLD = Font(bold=True)


def _num(v):
    return float(v) if v is not None and pd.notna(v) else None


def _calendar_values(wide, regions, months, replicate_from=None):
    """wide(달력월 × 지역) → {(region, month): 값}. replicate_from 열(예: '전국'·'서울')만 있으면 모든 지역에 복제"""
    out = {}
    if wide is None or wide.empty:
        return out
    cols = [c for c in wide.columns if c in regions]
    if not cols and replicate_from in wide.columns:
        s = wide[replicate_from]
        for r in regions:
            for m in months:
                out[(r, m)] = _num(s.get(m))
        return out
    for r in cols:
        s = wide[r]
        for m in months:
            out[(r, m)] = _num(s.get(m))
    return out


def _stepfill_monthly(wide, kind, months):
    """저빈도(연·분기·반기) → 기간 끝 달에 두고 다음 공표까지 유지(계단). 공표 시차는 수식 쪽에서 민다."""
    if kind == "M":
        return wide
    end = []
    for p in wide.index:
        s = str(p)
        if kind == "Y":
            end.append(pd.Period(f"{s[:4]}-12", "M"))
        elif kind == "Q":
            q = int(s[-1]) if "Q" in s else 4
            end.append(pd.Period(f"{s[:4]}-{q * 3:02d}", "M"))
        else:                                               # 반기 'YYYYH1/H2'
            end.append(pd.Period(f"{s[:4]}-{'06' if s.endswith('1') else '12'}", "M"))
    w = wide.set_axis(pd.PeriodIndex(end, freq="M")).sort_index()
    return w.reindex(pd.period_range(w.index.min(), months[-1], freq="M")).ffill()


# ------------------------------------------------------------ 수식
def _ok(i, lo, hi):
    return lo <= i <= hi


def formula_target(kind, i, n_first, n_last, n_regions, C="C", G3="G", G6="I"):
    if kind.startswith("past"):
        k = int(kind[4:])
        return f'=IF(AND(A{i - k}=A{i},ISNUMBER({C}{i}),ISNUMBER({C}{i - k})),100*({C}{i}/{C}{i - k}-1),"")' if _ok(i - k, n_first, n_last) else ""
    if kind in ("G3", "G6"):
        h = int(kind[1:])
        return f'=IF(AND(A{i + h}=A{i},ISNUMBER({C}{i + h}),ISNUMBER({C}{i})),100*({C}{i + h}/{C}{i}-1),"")' if _ok(i + h, n_first, n_last) else ""
    if kind in ("r3", "r6"):
        g = G3 if kind == "r3" else G6
        rng, brng = f"{g}${n_first}:{g}${n_last}", f"B${n_first}:B${n_last}"
        return f'=IF(AND(ISNUMBER({g}{i}),SUMPRODUCT(({brng}=B{i})*ISNUMBER({rng}))={n_regions}),{g}{i}-AVERAGEIFS({rng},{brng},B{i}),"")'
    return ""


def formula_x(key, i, X, lag, n_first, n_last, H=None, P=None, M=None):
    """설명변수 수식. X 원값 열, lag 공표 시차(행), H 세대수 보조 열, P 인구 보조 열, M 월값(누계 차분) 열. 미지원이면 None"""
    j = i - lag
    def ok(*rows):
        return all(_ok(r, n_first, n_last) for r in rows)
    def same(r):
        return f"A{r}=A{i}"
    if key in ("수준", "수준(0/±1)"):
        return f'=IF(AND({same(j)},ISNUMBER({X}{j})),{X}{j},"")' if ok(j) else ""
    if key == "수준_천명당":
        return f'=IF(AND({same(j)},ISNUMBER({X}{j}),ISNUMBER({P}{j})),1000*{X}{j}/{P}{j},"")' if (ok(j) and P) else ""
    if key == "수준_천세대당":
        return f'=IF(AND({same(j)},ISNUMBER({X}{j}),ISNUMBER({H}{j})),1000*{X}{j}/{H}{j},"")' if (ok(j) and H) else ""
    m = re.fullmatch(r"Δ(\d+)", key)
    if m:
        k = int(m.group(1))
        return f'=IF(AND({same(j - k)},ISNUMBER({X}{j}),ISNUMBER({X}{j - k})),{X}{j}-{X}{j - k},"")' if ok(j, j - k) else ""
    m = re.fullmatch(r"Δ(\d+)_천세대당", key)
    if m:
        k = int(m.group(1))
        return (f'=IF(AND({same(j - k)},ISNUMBER({X}{j}),ISNUMBER({X}{j - k}),ISNUMBER({H}{j}),ISNUMBER({H}{j - k})),1000*{X}{j}/{H}{j}-1000*{X}{j - k}/{H}{j - k},"")'
                if (ok(j, j - k) and H) else "")
    m = re.fullmatch(r"100·로그Δ(\d+)", key)
    if m:
        k = int(m.group(1))
        return f'=IF(AND({same(j - k)},ISNUMBER({X}{j}),ISNUMBER({X}{j - k}),{X}{j}>0,{X}{j - k}>0),100*(LN({X}{j})-LN({X}{j - k})),"")' if ok(j, j - k) else ""
    if key == "12개월합_천세대당":
        src = M or X
        return f'=IF(AND({same(j - 11)},COUNT({src}{j - 11}:{src}{j})=12,ISNUMBER({H}{j})),1000*SUM({src}{j - 11}:{src}{j})/{H}{j},"")' if (ok(j, j - 11) and H) else ""
    if key == "log1p(12개월합_천세대당)":
        src = M or X
        return f'=IF(AND({same(j - 11)},COUNT({src}{j - 11}:{src}{j})=12,ISNUMBER({H}{j})),LN(1+1000*SUM({src}{j - 11}:{src}{j})/{H}{j}),"")' if (ok(j, j - 11) and H) else ""
    if key == "12개월합 전년비(%)":
        src = M or X
        cur, prv = f"SUM({src}{j - 11}:{src}{j})", f"SUM({src}{j - 23}:{src}{j - 12})"
        return f'=IF(AND({same(j - 23)},COUNT({src}{j - 11}:{src}{j})=12,COUNT({src}{j - 23}:{src}{j - 12})=12,{prv}>0),100*({cur}/{prv}-1),"")' if ok(j, j - 23) else ""
    if key == "3개월합_천세대당":
        src = M or X
        return f'=IF(AND({same(j - 2)},COUNT({src}{j - 2}:{src}{j})=3,ISNUMBER({H}{j})),1000*SUM({src}{j - 2}:{src}{j})/{H}{j},"")' if (ok(j, j - 2) and H) else ""
    if key == "월값_천세대당":
        src = M or X
        return f'=IF(AND({same(j)},ISNUMBER({src}{j}),ISNUMBER({H}{j})),1000*{src}{j}/{H}{j},"")' if (ok(j) and H) else ""
    if key.startswith("log1p(천세대당"):
        return f'=IF(AND({same(j)},ISNUMBER({X}{j}),ISNUMBER({H}{j})),LN(1+1000*MAX({X}{j},0)/{H}{j}),"")' if (ok(j) and H) else ""
    m = re.fullmatch(r"(\d+)개월합_천명당", key)
    if m:
        k = int(m.group(1))
        return f'=IF(AND({same(j - k + 1)},COUNT({X}{j - k + 1}:{X}{j})={k},ISNUMBER({P}{j})),1000*SUM({X}{j - k + 1}:{X}{j})/{P}{j},"")' if (ok(j, j - k + 1) and P) else ""
    if key == "월값(누계":
        return f'=IF(RIGHT(B{i},2)="01",{X}{i},IF(AND({same(i - 1)},ISNUMBER({X}{i}),ISNUMBER({X}{i - 1})),{X}{i}-{X}{i - 1},""))' if ok(i - 1) else ""
    return None


FORMULA_DESC = {
    "R_last": "원값: 결정월 t 에 아는 마지막 월세통합가격지수 R(t−1) (raw V001, 공표시차 1개월)",
    "past1": "100·[R(t−1)/R(t−2) − 1] = 100*(C_t/C_{t−1} − 1)", "past3": "100*(C_t/C_{t−3} − 1)", "past6": "100*(C_t/C_{t−6} − 1)",
    "G3": "앞으로 3개월 변화율 100·[R(t+2)/R(t−1) − 1] = 100*(C_{t+3}/C_t − 1)", "G6": "100*(C_{t+6}/C_t − 1)",
    "r3": "G3 − 같은 결정월 패널 평균 (모든 지역이 있을 때만)", "r6": "G6 − 같은 결정월 패널 평균 (모든 지역이 있을 때만)",
}


# ------------------------------------------------------------ 시트 작성
def _write_target(wb, title, regions, R, months, desc):
    rows = [(r, m) for r in regions for m in months]
    n_first, n_last = 2, len(rows) + 1
    ws = wb.create_sheet(title)
    head = ["region", "결정월", "R_last", "past1", "past3", "past6", "G3", "r3", "G6", "r6", "구간"]
    ws.append(head)
    for c in range(1, len(head) + 1):
        ws.cell(row=1, column=c).font = BOLD
    for i, (reg, m) in enumerate(rows, start=2):
        rl = R.at[m - 1, reg] if ((m - 1) in R.index and reg in R.columns) else None
        ws.cell(row=i, column=1, value=reg)
        ws.cell(row=i, column=2, value=str(m))
        ws.cell(row=i, column=3, value=_num(rl))
        for col, kind in zip(range(4, 11), ["past1", "past3", "past6", "G3", "r3", "G6", "r6"]):
            f = formula_target(kind, i, n_first, n_last, len(regions))
            ws.cell(row=i, column=col, value=f if f else None)
        ws.cell(row=i, column=11, value="분석" if m >= ANALYSIS_START else "버퍼")
        if m < ANALYSIS_START:
            ws.cell(row=i, column=11).fill = GREY
    ws.freeze_panes = "C2"
    for k, v in FORMULA_DESC.items():
        desc.append({"시트": title, "열": k, "뜻": v})
    return head[2:10]


def _write_x(wb, title, regions, months, dec, vc, ctx, level, desc):
    rows = [(r, m) for r in regions for m in months]
    n_first, n_last = 2, len(rows) + 1
    wx = wb.create_sheet(title)
    wx.append(["region", "결정월", "구간"])
    for c in range(1, 4):
        wx.cell(row=1, column=c).font = BOLD
    for i, (reg, m) in enumerate(rows, start=2):
        wx.cell(row=i, column=1, value=reg)
        wx.cell(row=i, column=2, value=str(m))
        wx.cell(row=i, column=3, value="분석" if m >= ANALYSIS_START else "버퍼")
        if m < ANALYSIS_START:
            wx.cell(row=i, column=3).fill = GREY
    state = {"col": 4}
    helper_cols = {}
    rep = "전국" if level == "sido" else "서울"

    def add_values(name, values):
        c = state["col"]
        wx.cell(row=1, column=c, value=name).font = BOLD
        for i, key in enumerate(rows, start=2):
            wx.cell(row=i, column=c, value=values.get(key))
        state["col"] += 1
        return L(c)

    def add_formulas(name, maker):
        c = state["col"]
        wx.cell(row=1, column=c, value=name).font = BOLD
        for i in range(2, n_last + 1):
            f = maker(i)
            wx.cell(row=i, column=c, value=f if f else None)
        state["col"] += 1
        return L(c)

    def helper(which):
        if which in helper_cols:
            return helper_cols[which]
        existing = {str(wx.cell(row=1, column=c).value): L(c) for c in range(4, state["col"])}
        own = "V010_주민등록세대수_raw" if which == "hh" else "V009_주민등록인구_raw"
        if own in existing:                                   # 이미 결정된 변수의 raw 열이 있으면 그 열을 분모로 재사용
            helper_cols[which] = existing[own]
            return helper_cols[which]
        if level == "sido":
            src, label = (ctx["hh"], "V010_주민등록세대수_raw(보조)") if which == "hh" else (ctx["pop"], "V009_주민등록인구_raw(보조)")
        else:
            src, label = (ctx.get("hh_gu", pd.DataFrame()), "V010_주민등록세대수_raw(보조,구)") if which == "hh" else (ctx.get("pop_gu", pd.DataFrame()), "V009_주민등록인구_raw(보조,구)")
        helper_cols[which] = add_values(label, _calendar_values(src, regions, months))
        desc.append({"시트": title, "열": label, "뜻": "천세대당·천명당 계산용 보조 원값(달력월)"})
        return helper_cols[which]

    added = []
    for _, row in dec.iterrows():
        vid = row["ID"]
        name = re.sub(r"\s+", "", str(row["변수명"]))[:20]
        files = vc.find_raw_files(vid)
        if not files:
            continue
        parsed = [vc.parse_raw_file(f, vid) for f in files]
        replicated, derived_desc = False, ""
        der = vc.derived_raw(vid, parsed, level) or (vc.derived_raw(vid, parsed, "sido") if level == "gu" else None)
        if der is not None:                                     # 파생 집계 규칙이 있는 변수 (예: V011 비중)
            wide, dlabel, derived_desc = der
            name = f"{name}_{dlabel}"
            if level == "gu" and not any(c in regions for c in wide.columns):
                replicated = True
        else:
            items_sido = vc.select_items(vid, parsed, "sido")
            force_rep = level == "gu" and "서울 값 복제" in str(row.get("지역단위", ""))    # 결정 로그 지역단위에 '→ 서울 값 복제' 라고 적은 변수만 ('복제 불필요' 는 해당 없음)
            items = items_sido if (level == "sido" or force_rep) else vc.select_items(vid, parsed, "gu")
            if force_rep and items_sido:
                replicated = True
            if not items:
                if level == "gu" and items_sido:
                    items, replicated = items_sido, True      # 구 자료 없음 → 서울 시도 값 복제
                else:
                    continue
            wide = items[list(items)[0]]
        kind = wide.attrs.get("kind", "M")                     # 주 항목이 든 표의 주기 (파일이 여럿이면 표마다 다름)
        key = str(row["변환"]).strip()
        cand_names = set(vc.candidates(vid, wide, kind, vc.vtype_of(vid), ctx))
        if key not in cand_names:
            key = key.split(" ")[0]
        lag = int(row["공표시차_개월"]) if str(row["공표시차_개월"]).lstrip("-").isdigit() else 0
        lag += int(row["추가시차"]) if str(row["추가시차"]).lstrip("-").isdigit() else 0
        lag = max(lag, 0)
        wm = _stepfill_monthly(wide, kind, months)
        nat_only = list(wide.columns) == ["전국"]
        if kind != "M":
            # 저빈도: 변환(수준·전기 대비 차분·로그변화·전국대비)은 원 주기에서 계산한 뒤 기간 끝 달에 두고 다음 공표까지 유지(계단),
            # 결정월 행에서는 공표 시차만큼 위 행을 참조(수식). 원값 계단 열도 참고용으로 둔다.
            raw_label = f"{vid}_{name}_raw(계단채움)" + ("(전국값복제)" if nat_only else ("(서울값복제)" if replicated else ""))
            add_values(raw_label, _calendar_values(wm, regions, months, replicate_from=("전국" if nat_only else ("서울" if replicated else None))))
            desc.append({"시트": title, "열": raw_label, "뜻": f"raw 원값({ {'Y': '연', 'Q': '분기', 'H': '반기'}.get(kind, '저빈도') }간). 기간 끝 달에 두고 다음 공표까지 유지"})
            cands_lf = vc.candidates(vid, wide, kind, vc.vtype_of(vid), ctx)
            tw = cands_lf.get(key)
            if tw is None:
                continue
            tlabel = f"{vid}_{name}_{key}(계단)"
            T_ = add_values(tlabel, _calendar_values(_stepfill_monthly(tw, kind, months), regions, months, replicate_from=("전국" if nat_only else ("서울" if replicated else None))))
            desc.append({"시트": title, "열": tlabel, "뜻": f"{key} 를 원 주기에서 계산해 기간 끝 달에 두고 다음 공표까지 유지 (시차 반영 전)"})
            flabel = f"{vid}_{name}_{key}"
            add_formulas(flabel, lambda i: formula_x("수준", i, T_, lag, n_first, n_last))
            desc.append({"시트": title, "열": flabel, "뜻": f"공표시차 {lag}개월: 결정월 t 행 = 계단 열의 {lag}행 위(달력월 t−{lag}) 값"})
            added.append(flabel)
            continue
        raw_label = f"{vid}_{name}_raw" + ("(계단채움)" if kind != "M" else "") + ("(누계)" if vid in vc.CUMULATIVE else "") + ("(전국값복제)" if nat_only else ("(서울값복제)" if replicated else ""))
        X = add_values(raw_label, _calendar_values(wm, regions, months, replicate_from=("전국" if nat_only else ("서울" if replicated else None))))
        desc.append({"시트": title, "열": raw_label, "뜻": (derived_desc + "; " if derived_desc else "") + "raw 원값(달력월 기준" + (", 저빈도는 기간 끝 달에 두고 다음 공표까지 유지" if kind != "M" else "") + (", 전국 값을 모든 지역에 복제" if nat_only else "") + (", 구 자료가 없어 서울 시도 값을 복제" if replicated else "") + ")"})
        M = None
        if vid in vc.ZERO_GAPS and "0" in str(row.get("결측처리", "")):
            # KOSIS 가 값 0인 행을 생략한 표: 그 지역의 첫~마지막 관측 사이 빈 셀은 0 (보완 소분류 합계로 확인). 범위 밖 빈 셀은 결측 유지
            nm = len(months)

            def _zero_fill(i, X0=X):
                bs = n_first + ((i - n_first) // nm) * nm
                be = bs + nm - 1
                return f'=IF(ISNUMBER({X0}{i}),{X0}{i},IF(AND(COUNT({X0}{bs}:{X0}{i})>0,COUNT({X0}{i}:{X0}{be})>0),0,""))'
            zlabel = f"{vid}_{name}_raw(빈셀0)"
            X = add_formulas(zlabel, _zero_fill)
            desc.append({"시트": title, "열": zlabel, "뜻": "raw 빈 셀 중 그 지역의 첫~마지막 관측 사이에 있는 것은 0 (KOSIS 가 0인 행을 생략; 보완 소분류 파일의 합계 항등식으로 확인). 시작 전·종료 후 빈 셀은 그대로 결측"})
        if vid in vc.CUMULATIVE:
            M = add_formulas(f"{vid}_{name}_월값(누계차분)", lambda i, Xc=X: formula_x("월값(누계", i, Xc, 0, n_first, n_last))
            desc.append({"시트": title, "열": f"{vid}_{name}_월값(누계차분)", "뜻": "1월은 누계 그대로, 2월부터 전월 누계와의 차 (빈 셀은 앞 열에서 이미 0)"})
        H = helper("hh") if "천세대당" in key else None
        P = helper("pop") if "천명당" in key else None
        flabel = f"{vid}_{name}_{key}"
        probe = formula_x(key, 20, X, lag, n_first, n_last, H=H, P=P, M=M)
        if probe is None:                                    # 수식 미지원 → 값 직접 기입
            cands = vc.candidates(vid, wide if kind == "M" else wm, "M", vc.vtype_of(vid), ctx)
            w = cands.get(key)
            vals = _calendar_values(w.shift(lag), regions, months, replicate_from=("전국" if nat_only else ("서울" if replicated else None))) if w is not None else {}
            add_values(flabel, vals)
            desc.append({"시트": title, "열": flabel, "뜻": f"{key} (공표시차 {lag}개월 반영) — 수식 미지원이라 값 직접 기입"})
        else:
            add_formulas(flabel, lambda i: formula_x(key, i, X, lag, n_first, n_last, H=H, P=P, M=M))
            desc.append({"시트": title, "열": flabel, "뜻": f"{key}, 공표시차 {lag}개월: 결정월 t 행의 값 = 원값 열의 {lag}행 위(달력월 t−{lag})에 변환 적용"})
        added.append(flabel)
    wx.freeze_panes = "D2"
    return added


def write_workbook(ctx, log, vc):
    """ctx: variable_card.load_context 결과, log: 결정 로그 DataFrame, vc: variable_card 모듈"""
    R = ctx["targets"]["R_panel"]
    last = R.dropna(how="all").index.max() + 1               # 마지막 결정월 = 마지막 공표월 + 1
    months = pd.period_range(BUFFER_START, last, freq="M")
    dec = log[(log["상태"] == "결정") & (log["ID"] != "V001")]
    wb = Workbook()
    wb.remove(wb.active)
    desc = [{"시트": "공통", "열": "행", "뜻": f"지역 × 결정월 {months[0]}~{months[-1]} ({len(months)}개월). 2016-01 이전은 버퍼(12개월 변화 계산용). 수식은 같은 지역 블록 안 행만 참조, 아니면 \"\""}]
    tcols = _write_target(wb, "타깃", list(vc.REGIONS), R, months, desc)
    xcols = _write_x(wb, "설명변수", list(vc.REGIONS), months, dec, vc, ctx, "sido", desc)
    gu_summary = ""
    if ctx.get("targets_gu"):
        Rg = ctx["targets_gu"]["R_panel"]
        _write_target(wb, "타깃_서울구", list(vc.GU_LIST), Rg, months, desc)
        _write_x(wb, "설명변수_서울구", list(vc.GU_LIST), months, dec, vc, ctx, "gu", desc)
        gu_summary = f"서울 {len(vc.GU_LIST)}개 구 보조 패널 포함"
    wd = wb.create_sheet("설명")
    wd.append(["시트", "열", "뜻"])
    for c in range(1, 4):
        wd.cell(row=1, column=c).font = BOLD
    for d in desc:
        wd.append([d["시트"], d["열"], d["뜻"]])
    wd.column_dimensions["C"].width = 110
    os.makedirs(os.path.dirname(SHEET_PATH), exist_ok=True)
    try:
        wb.save(SHEET_PATH)
    except PermissionError:
        # Excel 에 열려 있으면 잠김 → 옆에 '_새본' 으로 저장해 두고, 닫힌 뒤 --apply 로 다시 쓴다
        alt = SHEET_PATH.replace(".xlsx", "_새본.xlsx")
        wb.save(alt)
        print(f"[경고] {os.path.relpath(SHEET_PATH, BASE)} 가 Excel 에 열려 있어 쓰지 못함 → {os.path.relpath(alt, BASE)} 로 저장. Excel 을 닫고 'variable_card.py --apply <ID>' 를 다시 실행하면 본 파일이 갱신됨")
    return {"행": len(vc.REGIONS) * len(months), "결정월": f"{months[0]}~{months[-1]}", "타깃 열": tcols, "설명변수 열": xcols, "보조": gu_summary}
