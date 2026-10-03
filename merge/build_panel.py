# -*- coding: utf-8 -*-
"""
[병합] processed/ -> 데이터취합_<YYYYMMDD>.xlsx

variables.py(데이터사전 Variable_Master 52개 변수 대응표)를 따라 processed 파일에서 변수를 찾아
지역 x 월 패널로 펼친다. 시트 구성(와꾸)만 docs/데이터취합.xlsx 예시를 참고했다. 변수는 데이터사전에서만 가져온다.

  17시도_월별     17개 시도 x 월. 시도 단위 월별 변수 + 전국 공통 월별 변수
  서울25구_월별   서울 25개 구 x 월. 구 단위 월별 변수 + 전국 공통 월별 변수
  분기_전국       분기 변수 (전국 단일값)
  반기_서울구고용  V013 고용률·취업자의 서울 구 단위 (지역별고용조사 반기)
  연간_시도·서울구 연간 변수
  변수현황        컬럼 ↔ Master ID ↔ 출처 ↔ 유효행·기간
  읽는법

병합 원칙 (데이터사전 README·Research_Notes 준수):
  - 월별이 아닌 자료는 월별 시트에 복제하지 않는다. 기준기간에 값을 채우면
    실제 공표 전 월에 미래정보가 들어간다(분기·반기·연간은 별도 시트).
  - 시도 값을 구에 복제하지 않는다. 서울25구 시트에는 구 단위 값이 있는 변수만 넣는다.
  - 전국 단일값(금리·물가 등)은 모든 지역에 같은 값으로 붙는다(공통 거시변수).
  - 빈칸은 0이 아니다(미공표·미제공).
  - 중복 변수(status=중복)는 대표 변수 한 컬럼만 둔다.

사용법:  python merge/build_panel.py            (기간 2011-01 ~ 최신월)
         python merge/build_panel.py 2015-06    (시작월 지정)
"""

import os
import sys
from datetime import date

import pandas as pd

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
from common import REGIONS, SEOUL_GU  # noqa: E402
from variables import BY_ID, RELEASE, RELEASE_MONTHLY, VARIABLES, drivers  # noqa: E402

START_YM = sys.argv[1] if len(sys.argv) > 1 else "2011-01"
OUT = os.path.join(BASE, f"데이터취합_{date.today():%Y%m%d}.xlsx")

# 행정표준코드 (현행). region_code = 코드 + 0 채움 (양식과 동일한 8자리)
SIDO_CODE = {"서울": "11", "부산": "26", "대구": "27", "인천": "28", "광주": "29", "대전": "30",
             "울산": "31", "세종": "36", "경기": "41", "강원": "51", "충북": "43", "충남": "44",
             "전북": "52", "전남": "46", "경북": "47", "경남": "48", "제주": "50"}
GU_CODE = {name: code for code, name in SEOUL_GU.items()}
GU_ORDER = list(SEOUL_GU.values())

_cache = {}


def load(rel):
    if rel not in _cache:
        df = pd.read_csv(os.path.join(BASE, rel), encoding="utf-8-sig", dtype=str)
        df["자치구"] = df["자치구"].fillna("")
        _cache[rel] = df
    return _cache[rel]


def select(v):
    df = load(v["processed"])
    stats = v["stat"] if isinstance(v["stat"], list) else [v["stat"]]
    df = df[df["통계명"].isin(stats)]
    if v["items"]:
        df = df[df["항목명"].isin(v["items"])]
    return df


def column_name(v, stat, item, unit, multi_stat, multi_item):
    """V011_20-29세_명 처럼 'ID_이름_단위'. 단위가 지수면 생략."""
    parts = [v["id"]]
    base = stat.replace("_아파트", "")
    if multi_stat and multi_item:
        parts += [base, item]
    elif multi_stat:
        parts.append(base)
    elif multi_item:
        parts.append(item)
    else:
        parts.append(base)
    name = "_".join(p for p in parts if p)
    unit = (unit or "").strip()
    if unit and unit != "지수" and not name.endswith(unit):
        name += f"_{unit}"
    return name.replace(" ", "")


# 자동 생성 이름이 뜻을 잃는 경우만 고친다 (생성명 -> 최종명)
RENAME = {
    "V012_인구이동_명": "V012_순이동자수_명",
    "V021_주택건설실적_호": "V021_아파트_인허가_호",
    "V022_주택건설실적_호": "V022_아파트_착공_호",
    "V023_주택건설실적_호": "V023_아파트_준공_호",
    "V006_전체_건": "V006_전월세_전체_건",
    "V006_전세_건": "V006_전월세_전세_건",
    "V006_월세(보증부포함)_건": "V006_전월세_월세(보증부포함)_건",
    "V006_신규계약_건": "V006_전월세_신규계약_건",
    "V006_갱신계약_건": "V006_전월세_갱신계약_건",
    "V007_거래(해제제외)_건": "V007_매매_거래(해제제외)_건",
    "V007_해제_건": "V007_매매_해제_건",
    "V035_동(호)수": "V035_아파트거래_동호수_호",
    "V035_면적_천㎡": "V035_아파트거래_면적_천㎡",
    "V036_동(호)수_호수": "V036_아파트매매거래_동호수_호",
    "V036_면적_천㎡": "V036_아파트매매거래_면적_천㎡",
    "V069_예금취급기관_가계대출_예금취급기관_십억원": "V069_예금취급기관_가계대출_십억원",
    "V069_예금취급기관_주택관련대출_주택관련대출-예금취급기관_십억원": "V069_예금취급기관_주택관련대출_십억원",
    "V015_지역소득_1인당_천원": "V015_1인당GRDP_천원",
    "V017_지역소득_1인당_천원": "V017_1인당가계총처분가능소득_천원",
    "V024_동수_동": "V024_주거용신축허가_동수_동",
    "V024_연면적_㎡": "V024_주거용신축허가_연면적_㎡",
    "V074_가계동향_도시2인이상_소득_원": "V074_가계동향_도시2인이상_월평균소득_원",
    "V077_임대주택건설공급현황_호": "V077_임대주택건설공급_총계_호",
    "V013H_취업자_천명": "V013_서울구반기_취업자_천명",
    "V013H_고용률_%": "V013_서울구반기_고용률_%",
}


def split_columns(v, df):
    """(통계명, 항목명) 조합마다 한 컬럼. -> [(컬럼명, 해당행 df)]"""
    combos = df[["통계명", "항목명"]].drop_duplicates().values.tolist()
    multi_stat = df["통계명"].nunique() > 1
    multi_item = df["항목명"].nunique() > 1
    out = []
    for stat, item in combos:
        part = df[(df["통계명"] == stat) & (df["항목명"] == item)]
        unit = part["단위"].dropna().iloc[0] if part["단위"].notna().any() else ""
        name = column_name(v, stat, item, unit, multi_stat, multi_item)
        out.append((RENAME.get(name, name), part))
    return out


def to_num(s):
    return pd.to_numeric(s, errors="coerce")


def monthly_frames():
    """월별 변수를 (시도 wide, 서울구 wide, 전국 wide) 로 펼친다."""
    sido, gu, nation, meta = {}, {}, {}, []
    for v in VARIABLES:
        if v["status"] != "수집" or v["freq"] != "M":
            continue
        df = select(v)
        for col, part in split_columns(v, df):
            part = part.assign(값=to_num(part["값"]))
            if v["level"] == "전국":
                s = part.set_index("연월")["값"]
                nation[col] = s[~s.index.duplicated()]
                where = "전국 공통"
            else:
                sd = part[part["자치구"] == ""]
                gd = part[part["자치구"] != ""]
                if not sd.empty:
                    sido[col] = sd.set_index(["광역지자체", "연월"])["값"]
                if not gd.empty:
                    gu[col] = gd.set_index(["자치구", "연월"])["값"]
                where = " / ".join(x for x, d in (("시도", sd), ("서울구", gd)) if not d.empty)
            meta.append(dict(컬럼=col, id=v["id"], where=where))
    return sido, gu, nation, meta


def build_monthly(frames, regions, key, panel, code_of, months, nation):
    idx = pd.MultiIndex.from_product([regions, months], names=[key, "연월"])
    out = pd.DataFrame(index=idx)
    for col, s in frames.items():
        out[col] = s.reindex(idx)
    for col, s in nation.items():
        out[col] = s.reindex(out.index.get_level_values("연월")).to_numpy()
    out = out.reset_index()
    out.insert(0, "panel", panel)
    out = out.rename(columns={key: "region"})
    out.insert(2, "region_code", out["region"].map(code_of))
    out.insert(3, "month", out.pop("연월").str.replace("-", "").astype(int))
    return out


def lowfreq_panel(freq):
    """분기·연간 변수를 (지역, 기간) wide 로."""
    parts, meta = [], []
    for v in VARIABLES:
        if v["status"] != "수집" or v["freq"] != freq:
            continue
        df = select(v)
        for col, part in split_columns(v, df):
            s = part.assign(값=to_num(part["값"])).set_index(["광역지자체", "자치구", "기간"])["값"]
            parts.append(s.rename(col))
            meta.append(dict(컬럼=col, id=v["id"], where="전국" if v["level"] == "전국" else v["level"]))
    wide = pd.concat(parts, axis=1).reset_index()
    order = {r: i for i, r in enumerate(["전국"] + REGIONS)}
    wide["_s"] = wide["광역지자체"].map(order).fillna(99)
    wide["_g"] = wide["자치구"].map(lambda g: -1 if g == "" else GU_ORDER.index(g) if g in GU_ORDER else 99)
    wide = wide.sort_values(["_s", "_g", "기간"]).drop(columns=["_s", "_g"])
    return wide.rename(columns={"광역지자체": "시도", "자치구": "구"}), meta


def coverage(df, cols, key):
    rows = {}
    for c in cols:
        s = df[[key, "month", c]].dropna(subset=[c])
        rows[c] = (len(s), f"{s['month'].min()}~{s['month'].max()}" if len(s) else "")
    return rows


def main():
    sido_f, gu_f, nation, meta_m = monthly_frames()

    all_months = set()
    for s in list(sido_f.values()) + list(gu_f.values()):
        all_months |= set(s.index.get_level_values("연월"))
    last = max(all_months)
    months = [f"{p.year:04d}-{p.month:02d}" for p in pd.period_range(START_YM, last, freq="M")]

    sido = build_monthly(sido_f, REGIONS, "광역지자체", "17개 시도",
                         lambda r: f"{SIDO_CODE[r]}000000", months, nation)
    gu = build_monthly(gu_f, GU_ORDER, "자치구", "서울 25개 구",
                       lambda r: f"{GU_CODE[r]}000", months, nation)
    quarterly, meta_q = lowfreq_panel("Q")
    annual, meta_y = lowfreq_panel("Y")
    halfyear, meta_h = lowfreq_panel("H")

    # ---------------------------------------------------------------- 변수현황
    cov_s = coverage(sido, sido.columns[4:], "region")
    cov_g = coverage(gu, gu.columns[4:], "region")
    rows = []
    for m in meta_m:
        v = BY_ID[m["id"]]
        ns, ps = cov_s.get(m["컬럼"], (0, ""))
        ng, pg = cov_g.get(m["컬럼"], (0, ""))
        rows.append([m["id"], m["컬럼"], v["name"], "월", m["where"],
                     "O" if m["컬럼"] in sido.columns else "", ns if m["컬럼"] in sido.columns else "",
                     ps if m["컬럼"] in sido.columns else "",
                     "O" if m["컬럼"] in gu.columns else "", ng if m["컬럼"] in gu.columns else "",
                     pg if m["컬럼"] in gu.columns else "",
                     v["source"], v["table"], v["note"]])
    for sheet, frame, meta, freq in (("분기_전국", quarterly, meta_q, "분기"),
                                     ("반기_서울구고용", halfyear, meta_h, "반기"),
                                     ("연간_시도·서울구", annual, meta_y, "연")):
        for m in meta:
            v = BY_ID[m["id"]]
            s = frame[["기간", m["컬럼"]]].dropna()
            rows.append([m["id"], m["컬럼"], v["name"], freq, f"{m['where']} ({sheet} 시트)",
                         "", "", "", "", "", "", v["source"], v["table"],
                         f"{len(s)}행 {s['기간'].min()}~{s['기간'].max()}. {v['note']}".strip()])
    for v in VARIABLES:
        if v["status"] == "중복":
            rows.append([v["id"], f"-> {v['same_as']}", v["name"], "", "대표 변수 컬럼 사용",
                         "", "", "", "", "", "", v["source"], v["table"], "중복 원자료"])
    status = pd.DataFrame(rows, columns=[
        "ID", "컬럼", "데이터사전 변수명", "주기", "지역 단위",
        "17시도", "17시도_유효행", "17시도_기간", "서울구", "서울구_유효행", "서울구_기간",
        "출처", "표ID/코드", "비고"])
    # 공표 시기 (공식 근거로 검증한 규칙. variables.RELEASE / RELEASE_MONTHLY)
    def release_cols(row):
        vid, freq = row["ID"], row["주기"]
        r = (RELEASE_MONTHLY.get(vid) if freq == "월" else None) or RELEASE.get(vid)             or (RELEASE.get("V013H") if vid == "V013H" else None)
        if r is None:
            src = BY_ID.get(vid, {}).get("same_as")
            return pd.Series(["", f"{src} 과 같음" if src else "", ""])
        lag = r["lag"]
        unit = "개월" if freq == "월" else "개월(기준기간 종료 후)"
        rule = (f"기준기간 내(전망 조사)" if lag < 0 else f"{lag}{unit}") + ((", 말일" if r["day"] >= 28 and r["day"] != 28 else f", {r['day']}일경") if r.get("day") else "")
        return pd.Series([rule, r["evidence"], r["source"]])
    status[["공표 시차", "공표 시기 (공식 근거)", "근거 출처"]] = status.apply(release_cols, axis=1)
    drv = drivers()
    i = status.columns.get_loc("데이터사전 변수명")
    status.insert(i, "동인", status["ID"].map(lambda x: drv.get(x, ("", ""))[0]))
    status["_o"] = status["ID"].str[1:4].astype(int)
    status = status.sort_values(["_o"], kind="stable").drop(columns="_o")

    readme = pd.DataFrame({"아파트 월세 연구 | 데이터사전(Variable_Master 52개) 변수 취합본": [
        f"기간 {months[0]} ~ {months[-1]} | 17개 시도 {len(sido):,}행 + 서울 25개 구 {len(gu):,}행 | 생성 {date.today()}",
        "",
        "먼저 확인할 것",
        "1. 변수는 docs/260928_데이터사전_취합_수정.xlsx 의 Variable_Master 시트(52개)만 담았다. 컬럼명 앞 V### 이 Master ID이고, 대응은 「변수현황」 시트와 variables.py.",
        "2. 빈칸은 0이 아니다. 미공표·미제공·해당 지역 없음이다.",
        "3. 17개 시도와 서울 25개 구는 서로 다른 패널이다. 서울시 행과 25개 구를 한 집합으로 중복 학습하지 않는다.",
        "4. 시도 단위 값을 구에 복제하지 않았다. 서울25구 시트에는 구 단위로 공표되는 변수만 있다.",
        "5. 전국 단일값(금리·물가·CSI 등)은 모든 지역 행에 같은 값이 들어간다. 주택가격전망CSI는 서울/6대광역시/기타도시 권역값 3개를 모두 둔다.",
        "6. 분기·반기·연간 자료는 월별 시트에 채우지 않았다(공표 전 월에 미래정보가 섞임). 별도 시트의 기준기간과 공표시차를 보고 붙인다.",
        "7. 목표변수 V001 월세통합가격지수는 2015.06부터다. 동일월 설명변수로 쓰지 않는다(leakage).",
        "8. V033 월세수급동향지수는 V005와 같은 통계라 V005 컬럼 하나로 둔다. V034 구 수급동향(2010.06~2015.06)은 현행과 별개 조사라 잇지 않는다.",
        "9. 광주·전남은 2026-07부터 통합 공표된다. 주민등록 계열은 광주 5개 구 합계로 광주, 나머지로 전남을 복원했다.",
        "10. V006·V007 실거래 건수는 신고 원본 행수(월별 최신 확인분)이며 공식 거래량이 아니다. 최근 2~3개월은 신고·해제로 계속 바뀐다.",
        "11. V074 가계동향 소득은 사전 수록기간대로 2019Q4에서 끝나는 구계열이다(2019년 표본개편 이후 신계열은 사전에 없음).",
        "",
        "시트 구성",
        "17시도_월별 : 시도 x 월. 시도 단위 월별 변수 + 전국 공통 월별 변수",
        "서울25구_월별 : 서울 구 x 월. 구 단위 월별 변수 + 전국 공통 월별 변수",
        "분기_전국 : 가계신용·대출태도·시간당 명목임금지수·가계동향 소득 (기간 = 2026Q2 형식)",
        "반기_서울구고용 : V013 고용률·취업자의 서울 구 단위 (2021H1~, 기간 = 2021H1 형식)",
        "연간_시도·서울구 : 1인당 GRDP·가계총처분가능소득, 서울 구 주거용 신축허가, 주택보급률(시도), 임대주택",
        "변수현황 : 컬럼별 Master ID·출처·유효행·기간",
    ]})

    # 열 제목을 "ID_무슨 통계_세부_단위" 로 (전처리 파일과 같은 이름표: impute/display_names.py)
    import re
    sys.path.insert(0, os.path.join(BASE, "impute"))
    from display_names import DISPLAY, disp
    sido, gu, quarterly, halfyear, annual = [df.rename(columns=disp) for df in (sido, gu, quarterly, halfyear, annual)]
    names = sorted((c for c in DISPLAY if DISPLAY[c] != c), key=len, reverse=True)
    pat = re.compile("|".join(re.escape(n) for n in names))
    fix = lambda v: pat.sub(lambda m: DISPLAY[m.group(0)], v) if isinstance(v, str) else v  # noqa: E731
    readme, status = [df.apply(lambda s: s.map(fix)) for df in (readme, status)]

    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        readme.to_excel(xw, sheet_name="읽는법", index=False)
        status.to_excel(xw, sheet_name="변수현황", index=False)
        sido.to_excel(xw, sheet_name="17시도_월별", index=False)
        gu.to_excel(xw, sheet_name="서울25구_월별", index=False)
        quarterly.to_excel(xw, sheet_name="분기_전국", index=False)
        halfyear.to_excel(xw, sheet_name="반기_서울구고용", index=False)
        annual.to_excel(xw, sheet_name="연간_시도·서울구", index=False)
        style(xw.book)

    print(f"완료: {OUT}")
    print(f"  17시도_월별  {sido.shape[0]:,}행 x {sido.shape[1]}열")
    print(f"  서울25구_월별 {gu.shape[0]:,}행 x {gu.shape[1]}열")
    print(f"  분기_전국 {quarterly.shape} / 반기_서울구고용 {halfyear.shape} / 연간 {annual.shape}")


def style(wb):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head = PatternFill("solid", fgColor="DDE6F0")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=10, bold=(c.row == 1))
        if ws.title == "읽는법":
            ws.column_dimensions["A"].width = 140
            continue
        for c in ws[1]:
            c.fill = head
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 45
        ws.freeze_panes = "E2" if ws.title.endswith("월별") else "B2"
        for i, c in enumerate(ws[1], 1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(c.value or "")) * 1.3, 9), 28)
        ws.auto_filter.ref = ws.dimensions


if __name__ == "__main__":
    main()
